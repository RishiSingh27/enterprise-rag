"""
Layout-aware document parser (T022-T031).

Constitution Article III: naive `pdf.extract_text()` is PROHIBITED. We need
structure — headings, reading order, table cells, figure locations — because
structure is what makes chunking and retrieval correct.

Two backends, same output contract (Element list):
  * azure  - Azure AI Document Intelligence prebuilt-layout (default)
  * local  - pdfplumber + python-docx/pptx/openpyxl + Tesseract (offline)

Selected by PARSER_BACKEND in .env, so the pipeline never changes.
"""

from __future__ import annotations

import logging
from pathlib import Path

from .chunker import Element

log = logging.getLogger("ekb.parser")

SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".png", ".jpg", ".jpeg"}


class UnsupportedFileType(Exception):
    pass


class ParseError(Exception):
    pass


class DocumentParser:
    """Facade over the two backends."""

    def __init__(self, settings, multimodal=None):
        self.settings = settings
        self.multimodal = multimodal

    # ------------------------------------------------------------------ #
    # Validation (FR-001)
    # ------------------------------------------------------------------ #

    def validate(self, file_path: Path, file_bytes: bytes) -> None:
        suffix = file_path.suffix.lower()
        if suffix not in SUPPORTED:
            raise UnsupportedFileType(
                f"Unsupported file type '{suffix}'. Supported: {', '.join(sorted(SUPPORTED))}"
            )
        if not file_bytes:
            raise ParseError("File is empty (0 bytes).")
        if len(file_bytes) > self.settings.max_upload_mb * 1024 * 1024:
            raise ParseError(f"File exceeds {self.settings.max_upload_mb} MB limit.")
        if suffix == ".pdf" and file_bytes[:5] != b"%PDF-":
            raise ParseError("File is not a valid PDF (missing %PDF- header).")

    # ------------------------------------------------------------------ #
    # Stage 1 - cheap text pass, used ONLY for Tier-2 page hashing
    # ------------------------------------------------------------------ #

    def extract_page_text(self, file_path: Path) -> list[tuple[int, str]]:
        """
        Fast, cheap text-per-page extraction. This runs BEFORE we decide which
        pages changed, so it must never call a paid service — otherwise the
        whole cost-saving design is defeated at the first step.
        """
        suffix = file_path.suffix.lower()
        try:
            if suffix == ".pdf":
                import pdfplumber
                with pdfplumber.open(file_path) as pdf:
                    return [(i, page.extract_text() or "") for i, page in enumerate(pdf.pages, 1)]

            if suffix == ".docx":
                import docx
                doc = docx.Document(str(file_path))
                text = "\n".join(p.text for p in doc.paragraphs)
                return [(1, text)]

            if suffix == ".pptx":
                from pptx import Presentation
                prs = Presentation(str(file_path))
                pages = []
                for i, slide in enumerate(prs.slides, 1):
                    parts = [sh.text for sh in slide.shapes if hasattr(sh, "text")]
                    pages.append((i, "\n".join(parts)))
                return pages

            if suffix == ".xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True)
                pages = []
                for i, ws in enumerate(wb.worksheets, 1):
                    rows = [
                        " ".join(str(c) for c in row if c is not None)
                        for row in ws.iter_rows(values_only=True)
                    ]
                    pages.append((i, "\n".join(rows)))
                return pages

            if suffix in (".txt", ".md"):
                return [(1, file_path.read_text(encoding="utf-8", errors="replace"))]

            if suffix in (".png", ".jpg", ".jpeg"):
                return [(1, "")]      # image-only; hash relies on structure signature

        except Exception as exc:  # noqa: BLE001
            raise ParseError(f"Could not read {file_path.name}: {exc}") from exc

        raise UnsupportedFileType(f"No text extractor for '{suffix}'")

    def quick_structure(self, file_path: Path) -> dict[int, list[dict]]:
        """
        Cheap structural fingerprint per page for the Tier-2 hash (FR-004).
        Counts tables and images without running the full layout model.
        """
        suffix = file_path.suffix.lower()
        result: dict[int, list[dict]] = {}
        try:
            if suffix == ".pdf":
                import pdfplumber
                with pdfplumber.open(file_path) as pdf:
                    for i, page in enumerate(pdf.pages, 1):
                        elements = []
                        order = 0
                        for line in (page.extract_text() or "").splitlines():
                            stripped = line.strip()
                            if not stripped:
                                continue
                            kind = "heading" if _looks_like_heading(stripped) else "paragraph"
                            elements.append({"element_type": kind, "reading_order": order})
                            order += 1
                        for _ in page.find_tables():
                            elements.append({"element_type": "table", "reading_order": order})
                            order += 1
                        for _ in (page.images or []):
                            elements.append({"element_type": "image", "reading_order": order})
                            order += 1
                        result[i] = elements
                return result
        except Exception as exc:  # noqa: BLE001
            log.warning("quick_structure.failed", extra={"error": str(exc)[:200]})

        for page_number, text in self.extract_page_text(file_path):
            elements = []
            for order, line in enumerate(t for t in text.splitlines() if t.strip()):
                kind = "heading" if _looks_like_heading(line.strip()) else "paragraph"
                elements.append({"element_type": kind, "reading_order": order})
            result[page_number] = elements
        return result

    # ------------------------------------------------------------------ #
    # Stage 2 - full parse, CHANGED PAGES ONLY (FR-006, FR-011..FR-015)
    # ------------------------------------------------------------------ #

    def parse_pages(self, file_path: Path, page_numbers: list[int]) -> list[Element]:
        if not page_numbers:
            return []
        if self.settings.parser_backend == "azure":
            try:
                return self._parse_azure(file_path, page_numbers)
            except Exception as exc:  # noqa: BLE001
                log.warning("parser.azure_failed_falling_back", extra={"error": str(exc)[:200]})
        return self._parse_local(file_path, page_numbers)

    # ---------------------- Azure backend ---------------------- #

    def _parse_azure(self, file_path: Path, page_numbers: list[int]) -> list[Element]:
        from ..core.azure_clients import get_document_intelligence_client

        client = get_document_intelligence_client()
        if client is None:
            raise ParseError("Document Intelligence not configured")

        with open(file_path, "rb") as fh:
            poller = client.begin_analyze_document(
                "prebuilt-layout",
                body=fh,
                pages=",".join(str(n) for n in page_numbers),
                content_type="application/octet-stream",
            )
        result = poller.result()

        elements: list[Element] = []
        wanted = set(page_numbers)

        for para in (result.paragraphs or []):
            page = _page_of(para)
            if page not in wanted:
                continue
            role = (getattr(para, "role", None) or "").lower()
            if role in ("title", "sectionheading"):
                kind, level = "heading", (1 if role == "title" else 2)
            elif role in ("pagefooter", "pageheader", "pagenumber"):
                continue                      # chrome, never content
            else:
                kind, level = "paragraph", None
            elements.append(Element(
                element_type=kind, content=para.content, page_number=page,
                reading_order=len(elements), heading_level=level,
            ))

        # FR-012: tables as Markdown, header row preserved
        for table in (result.tables or []):
            page = _page_of(table)
            if page not in wanted:
                continue
            elements.append(Element(
                element_type="table",
                content=_table_to_markdown(table),
                page_number=page,
                reading_order=len(elements),
            ))

        # FR-013: figures described by the vision model
        if self.settings.enable_multimodal and self.multimodal:
            for figure in (getattr(result, "figures", None) or []):
                page = _page_of(figure)
                if page not in wanted:
                    continue
                described = self.multimodal.describe_figure(file_path, figure, page)
                if described:
                    description, image_path = described
                    elements.append(Element(
                        element_type="image", content=description, page_number=page,
                        reading_order=len(elements), image_path=image_path,
                    ))

        return elements

    # ---------------------- Local backend ---------------------- #

    def _parse_local(self, file_path: Path, page_numbers: list[int]) -> list[Element]:
        """Offline fallback (T030). Lower fidelity, but unblocks restricted networks."""
        suffix = file_path.suffix.lower()
        elements: list[Element] = []
        wanted = set(page_numbers)

        if suffix == ".pdf":
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    if i not in wanted:
                        continue
                    text = page.extract_text() or ""
                    if not text.strip() and self.settings.enable_ocr:
                        text = _ocr_page(page)
                    for line in text.splitlines():
                        stripped = line.strip()
                        if not stripped:
                            continue
                        is_heading = _looks_like_heading(stripped)
                        elements.append(Element(
                            element_type="heading" if is_heading else "paragraph",
                            content=stripped, page_number=i,
                            reading_order=len(elements),
                            heading_level=2 if is_heading else None,
                        ))
                    for table in page.extract_tables():
                        elements.append(Element(
                            element_type="table",
                            content=_rows_to_markdown(table),
                            page_number=i, reading_order=len(elements),
                        ))
            return elements

        for page_number, text in self.extract_page_text(file_path):
            if page_number not in wanted:
                continue
            for line in text.splitlines():
                stripped = line.strip()
                if not stripped:
                    continue
                is_heading = _looks_like_heading(stripped)
                elements.append(Element(
                    element_type="heading" if is_heading else "paragraph",
                    content=stripped, page_number=page_number,
                    reading_order=len(elements),
                    heading_level=2 if is_heading else None,
                ))
        return elements


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #

def _looks_like_heading(line: str) -> bool:
    """Heuristic for the local backend: short, title-ish, no terminal period."""
    if len(line) > 90 or line.endswith("."):
        return False
    if line.startswith("#"):
        return True
    words = line.split()
    if not 1 <= len(words) <= 12:
        return False
    return line.isupper() or line.istitle()


def _page_of(item) -> int:
    try:
        regions = getattr(item, "bounding_regions", None) or []
        if regions:
            return regions[0].page_number
        spans = getattr(item, "spans", None) or []
        if spans:
            return 1
    except Exception:  # noqa: BLE001
        pass
    return 1


def _table_to_markdown(table) -> str:
    """FR-012. Preserve the header row — without it the numbers lose meaning."""
    if not table.cells:
        return ""
    n_rows = table.row_count
    n_cols = table.column_count
    grid = [["" for _ in range(n_cols)] for _ in range(n_rows)]
    for cell in table.cells:
        r, c = cell.row_index, cell.column_index
        if r < n_rows and c < n_cols:
            grid[r][c] = (cell.content or "").replace("|", "\\|").replace("\n", " ").strip()
    return _rows_to_markdown(grid)


def _rows_to_markdown(rows) -> str:
    rows = [[("" if c is None else str(c)).replace("|", "\\|").replace("\n", " ").strip()
             for c in row] for row in rows if row]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    lines = ["| " + " | ".join(header) + " |",
             "|" + "|".join("---" for _ in header) + "|"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def _ocr_page(page) -> str:
    """FR-014. Scanned page with no text layer."""
    try:
        import pytesseract
        image = page.to_image(resolution=200).original
        return pytesseract.image_to_string(image)
    except Exception as exc:  # noqa: BLE001
        log.warning("ocr.failed", extra={"error": str(exc)[:200]})
        return ""
